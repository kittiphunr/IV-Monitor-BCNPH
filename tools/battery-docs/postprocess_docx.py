#!/usr/bin/env python3
"""ตกแต่ง .docx ที่ pandoc สร้าง: ตารางเต็มความกว้างหน้า มีเส้นขอบ และหัวตารางตัวหนา"""
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

path = sys.argv[1]
doc = Document(path)
for t in doc.tables:
    tblPr = t._tbl.tblPr
    w = tblPr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW"); tblPr.append(w)
    w.set(qn("w:type"), "pct"); w.set(qn("w:w"), "5000")
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
    lay.set(qn("w:type"), "autofit")
    b = tblPr.find(qn("w:tblBorders"))
    if b is not None:
        tblPr.remove(b)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "A0A0A0")
        b.append(e)
    tblPr.append(b)
    # ความกว้างคอลัมน์: แบ่งความกว้างหน้า 16 ซม. ตามความยาวข้อความยาวสุดของแต่ละคอลัมน์
    # (รากที่สอง ให้คอลัมน์สั้นไม่แคบเกินไป) แล้วตรึงไว้ เพราะ LibreOffice ไม่ขยายตาราง autofit ให้เต็มหน้า
    TOTAL = 9072  # twips = 16 cm
    ncol = len(t.columns)
    longest = [1] * ncol
    for row in t.rows:
        for i, c in enumerate(row.cells[:ncol]):
            longest[i] = max(longest[i], len(c.text))
    weight = [max(n, 4) ** 0.5 for n in longest]
    widths = [int(TOTAL * x / sum(weight)) for x in weight]
    lay.set(qn("w:type"), "fixed")
    grid = t._tbl.tblGrid
    for gc, wd in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(wd))
    for row in t.rows:
        for i, c in enumerate(row.cells[:ncol]):
            tcPr = c._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
            tcW.set(qn("w:type"), "dxa"); tcW.set(qn("w:w"), str(widths[i]))
    for c in t.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
doc.save(path)
